"""Export DocSpec JSON to DOCX via python-docx."""

from __future__ import annotations

from pathlib import Path
from typing import Any


def _add_paragraphs(doc, text: str) -> None:
    for block in (text or "").split("\n\n"):
        block = block.strip()
        if not block:
            continue
        if block.startswith("# "):
            doc.add_heading(block[2:].strip(), level=1)
        elif block.startswith("## "):
            doc.add_heading(block[3:].strip(), level=2)
        elif block.startswith("### "):
            doc.add_heading(block[4:].strip(), level=3)
        else:
            for line in block.splitlines():
                line = line.strip()
                if line.startswith(("- ", "* ")):
                    doc.add_paragraph(line[2:].strip(), style="List Bullet")
                elif line[:3].rstrip(".").isdigit() and ". " in line[:6]:
                    doc.add_paragraph(line.split(". ", 1)[-1], style="List Number")
                else:
                    doc.add_paragraph(line)


def export_docx(spec: dict[str, Any], output_path: Path) -> None:
    from docx import Document
    from docx.shared import Pt

    doc = Document()
    title = (spec.get("title") or "Document").strip() or "Document"
    doc.add_heading(title, level=0)

    subtitle = (spec.get("subtitle") or "").strip()
    if subtitle:
        p = doc.add_paragraph(subtitle)
        for run in p.runs:
            run.italic = True
            run.font.size = Pt(11)

    sections = spec.get("sections") or []
    if not sections and spec.get("body_markdown"):
        sections = [{"heading": "", "body_markdown": spec.get("body_markdown"), "bullets": []}]

    for section in sections:
        heading = (section.get("heading") or "").strip()
        if heading:
            doc.add_heading(heading, level=1)
        body = (section.get("body_markdown") or section.get("body") or "").strip()
        if body:
            _add_paragraphs(doc, body)
        for bullet in section.get("bullets") or []:
            text = str(bullet).strip()
            if text:
                doc.add_paragraph(text, style="List Bullet")

        table = section.get("table")
        if isinstance(table, dict):
            headers = table.get("headers") or []
            rows = table.get("rows") or []
            if headers:
                t = doc.add_table(rows=1 + len(rows), cols=len(headers))
                t.style = "Table Grid"
                for i, h in enumerate(headers):
                    t.rows[0].cells[i].text = str(h)
                for r_i, row in enumerate(rows):
                    for c_i, cell in enumerate(row):
                        if c_i < len(headers):
                            t.rows[r_i + 1].cells[c_i].text = str(cell)

    # Top-level tables
    for table in spec.get("tables") or []:
        if not isinstance(table, dict):
            continue
        headers = table.get("headers") or []
        rows = table.get("rows") or []
        if not headers:
            continue
        t = doc.add_table(rows=1 + len(rows), cols=len(headers))
        t.style = "Table Grid"
        for i, h in enumerate(headers):
            t.rows[0].cells[i].text = str(h)
        for r_i, row in enumerate(rows):
            for c_i, cell in enumerate(row):
                if c_i < len(headers):
                    t.rows[r_i + 1].cells[c_i].text = str(cell)

    output_path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(output_path))
