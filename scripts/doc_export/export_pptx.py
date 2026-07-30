"""Export DocSpec JSON to PPTX via python-pptx."""

from __future__ import annotations

from pathlib import Path
from typing import Any


def export_pptx(spec: dict[str, Any], output_path: Path) -> None:
    from pptx import Presentation
    from pptx.util import Pt

    prs = Presentation()
    title = (spec.get("title") or "Presentation").strip() or "Presentation"
    subtitle = (spec.get("subtitle") or "").strip()

    # Title slide
    title_layout = prs.slide_layouts[0]
    slide = prs.slides.add_slide(title_layout)
    slide.shapes.title.text = title
    if len(slide.placeholders) > 1:
        slide.placeholders[1].text = subtitle

    slides = spec.get("slides") or []
    if not slides:
        # Derive from sections when the model returned a docx-like spec
        for section in spec.get("sections") or []:
            slides.append(
                {
                    "title": section.get("heading") or title,
                    "bullets": section.get("bullets")
                    or [
                        line.strip("-* ").strip()
                        for line in (section.get("body_markdown") or "").splitlines()
                        if line.strip()
                    ][:8],
                    "notes": "",
                }
            )

    bullet_layout = prs.slide_layouts[1]
    for item in slides:
        s = prs.slides.add_slide(bullet_layout)
        s.shapes.title.text = str(item.get("title") or "Slide").strip() or "Slide"
        body = s.shapes.placeholders[1].text_frame
        body.clear()
        bullets = item.get("bullets") or []
        if not bullets and item.get("body"):
            bullets = [str(item.get("body"))]
        if not bullets:
            bullets = [" "]
        for i, bullet in enumerate(bullets):
            text = str(bullet).strip()
            if not text:
                continue
            if i == 0:
                p = body.paragraphs[0]
            else:
                p = body.add_paragraph()
            p.text = text
            p.level = 0
            p.font.size = Pt(20)
        notes = (item.get("notes") or "").strip()
        if notes:
            s.notes_slide.notes_text_frame.text = notes

    output_path.parent.mkdir(parents=True, exist_ok=True)
    prs.save(str(output_path))
